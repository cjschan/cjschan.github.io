#!/usr/bin/env python3
"""Write the student-facing weekly to-do list (one section per week) from course_map_data.py.

Usage:  python3 build_weekly_todo.py
Output: ../Math1342_Weekly_ToDo.docx  (opens in Google Docs or Word; copy each week into Blackboard)
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "Math1342_Weekly_ToDo.docx")
sys.path.insert(0, HERE)
import course_map_data as D  # noqa: E402

def verb(text):
    """Student-facing action for a learning material."""
    if text.startswith(("e-Text", "START HERE")):
        return "Read"
    if text.startswith("Interactive Lecture"):
        return "Watch"
    if text.startswith("StatKey"):
        return "Use"
    if "Slides" in text.split(":")[0] or text.startswith("All "):
        return "Review"
    return None

ITEM_PREFIXES = ("e-Text", "Interactive Lecture", "Slides", "Online Homework", "Participation", "Take-Home Quiz",
                 "Test", "Practice Test", "Final Exam", "START HERE", "Introduction Slides", "Review Slides",
                 "StatKey", "All ", "Course Evaluation", "Your Turn")

def split_name(text):
    """Bold the Blackboard item name; keep due dates and descriptions in regular type."""
    if not text.startswith(ITEM_PREFIXES):
        return "", text
    m = re.search(r" \(| – ", text)
    return (text[:m.start()], text[m.start():]) if m else (text, "")

def clean(text):
    return re.sub(r" \((?:all objectives|objectives? [^()]+)\)$", "", text)

def items(w, kind):
    """One week's items of a kind, in row order, each listed once."""
    out = []
    for r in w["rows"]:
        for text in r[kind]:
            if text not in out:
                out.append(text)
    return out

def bullet(doc, text, prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if prefix:
        p.add_run(prefix + " ")
    name, rest = split_name(clean(text))
    if name:
        p.add_run(name).bold = True
    if rest:
        p.add_run(rest)
    p.paragraph_format.space_after = Pt(2)

def label(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)

def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    title = doc.add_heading(f"{D.COURSE}: What to Do Each Week", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(D.TERM)
    doc.add_paragraph("Each week has three parts: read and watch the materials, take part in the class "
                      "activities, and complete the assignments by their due dates. All times are 11:59 PM "
                      "unless noted.")
    # Quizzes and test items sit in every week whose objectives they cover; students see each once, in the last of those weeks.
    last_week = {clean(t): w["num"] for w in D.WEEKS for t in items(w, "assessments")}
    # A test that sits on an "also" line (the week it happens) is listed there instead of in its last row week.
    tests_in_also = {m.group(1) for w in D.WEEKS for a in w["also"] for m in [re.match(r"Test (\d+) –", a)] if m}
    for w in D.WEEKS:
        doc.add_heading(f"Week {w['num']}: {w['title']} ({w['dates']})", level=1)
        p = doc.add_paragraph()
        p.add_run("Topics: ").bold = True
        p.add_run("; ".join(w["sections"]))
        label(doc, "Read and watch")
        for text in items(w, "materials"):
            bullet(doc, text, verb(text))
        label(doc, "Participate")
        for text in items(w, "activities"):
            bullet(doc, text)
        label(doc, "Complete and submit")
        # Repeated items show only in their last week.
        work = [t for t in items(w, "assessments") if last_week[clean(t)] == w["num"]
                and not any(t.startswith(f"Test {n} (") for n in tests_in_also)]
        for text in work + [a for a in w["also"] if a.startswith(ITEM_PREFIXES)]:
            bullet(doc, text)
    doc.save(OUT)
    print(f"Wrote {OUT}")

if __name__ == "__main__":
    main()
